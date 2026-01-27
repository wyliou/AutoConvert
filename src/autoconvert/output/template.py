"""Template population for 40-column output."""

from pathlib import Path

from openpyxl import Workbook, load_workbook

from autoconvert.core.models import InvoiceItem, PackingTotals

# Fixed values for output columns
FIXED_EXEMPTION_METHOD = "3"
FIXED_DESTINATION_REGION = "32052"
FIXED_ADMIN_CODE = "320506"
FIXED_DESTINATION_COUNTRY = "142"

# Column indices (1-based) for output template
COL_PART_NO = 1  # A
COL_PO_NO = 2  # B
COL_EXEMPTION = 3  # C
COL_CURRENCY = 4  # D
COL_QTY = 5  # E
COL_PRICE = 6  # F
COL_AMOUNT = 7  # G
COL_COO = 8  # H
# I-K reserved
COL_SERIAL = 12  # L
COL_WEIGHT = 13  # M
COL_INV_NO = 14  # N
# O reserved
COL_TOTAL_GW = 16  # P (row 5 only)
# Q reserved
COL_DEST_REGION = 18  # R
COL_ADMIN_CODE = 19  # S
COL_DEST_COUNTRY = 20  # T
# U-AJ reserved
COL_TOTAL_PACKETS = 37  # AK (row 5 only)
COL_BRAND = 38  # AL
COL_BRAND_TYPE = 39  # AM
COL_MODEL = 40  # AN


def populate_template(
    template_path: Path,
    invoice_items: list[InvoiceItem],
    packing_totals: PackingTotals,
) -> Workbook:
    """Populate output template with invoice data.

    Args:
        template_path (Path): Path to output_template.xlsx.
        invoice_items (list[InvoiceItem]): Transformed invoice items with weights.
        packing_totals (PackingTotals): Packing totals for row 5.

    Returns:
        Workbook: Populated workbook ready for saving.
    """
    # Load fresh copy of template to preserve formatting
    wb = load_workbook(template_path)
    sheet = wb.active

    if sheet is None:
        raise ValueError("Template has no active sheet")

    # Data rows start at row 5
    start_row = 5

    for idx, item in enumerate(invoice_items):
        row = start_row + idx

        # Write item data
        sheet.cell(row=row, column=COL_PART_NO, value=item.part_no)
        sheet.cell(row=row, column=COL_PO_NO, value=item.po_no)
        sheet.cell(row=row, column=COL_EXEMPTION, value=FIXED_EXEMPTION_METHOD)
        sheet.cell(row=row, column=COL_CURRENCY, value=item.currency)
        sheet.cell(row=row, column=COL_QTY, value=item.qty)
        sheet.cell(row=row, column=COL_PRICE, value=item.price)
        sheet.cell(row=row, column=COL_AMOUNT, value=item.amount)
        sheet.cell(row=row, column=COL_COO, value=item.coo)
        sheet.cell(row=row, column=COL_SERIAL, value=item.serial if item.serial else None)
        sheet.cell(row=row, column=COL_WEIGHT, value=item.weight)
        sheet.cell(row=row, column=COL_INV_NO, value=item.inv_no)
        sheet.cell(row=row, column=COL_DEST_REGION, value=FIXED_DESTINATION_REGION)
        sheet.cell(row=row, column=COL_ADMIN_CODE, value=FIXED_ADMIN_CODE)
        sheet.cell(row=row, column=COL_DEST_COUNTRY, value=FIXED_DESTINATION_COUNTRY)
        sheet.cell(row=row, column=COL_BRAND, value=item.brand)
        sheet.cell(row=row, column=COL_BRAND_TYPE, value=item.brand_type)
        sheet.cell(row=row, column=COL_MODEL, value=item.model)

        # Row 5 special fields
        if idx == 0:
            sheet.cell(row=row, column=COL_TOTAL_GW, value=float(packing_totals.total_gw))
            if packing_totals.total_packets:
                sheet.cell(row=row, column=COL_TOTAL_PACKETS, value=packing_totals.total_packets)

    return wb
