"""Column mapping using regex patterns.

Maps vendor column headers to standardized field names using patterns
from field_patterns.yaml.
"""

import re
from typing import Literal

from openpyxl.worksheet.worksheet import Worksheet

from autoconvert.core.errors import ErrorCode
from autoconvert.core.models import ColumnMap, ValidationResult
from autoconvert.logging.messages import log_debug
from autoconvert.setup.registry import PatternRegistry

# Currency codes for data row detection
CURRENCY_CODES = {"USD", "CNY", "EUR", "RMB", "JPY", "GBP", "HKD", "TWD"}

# Keywords that indicate a price/amount column (for merged header detection)
PRICE_AMOUNT_KEYWORDS = {"PRICE", "AMOUNT", "金额", "单价"}


def normalize_header(value: str | None) -> str:
    """Normalize header text for pattern matching.

    - Collapses newlines, tabs, multiple spaces to single space
    - Strips leading/trailing whitespace

    Args:
        value (str | None): Raw header value.

    Returns:
        str: Normalized header string.
    """
    if value is None:
        return ""

    text = str(value)
    # Replace newlines and tabs with space
    text = re.sub(r"[\n\r\t]+", " ", text)
    # Collapse multiple spaces
    text = re.sub(r" +", " ", text)
    return text.strip()


def map_columns(
    sheet: Worksheet,
    header_row: int,
    registry: PatternRegistry,
    sheet_type: Literal["invoice", "packing"],
    validation: ValidationResult,
) -> ColumnMap | None:
    """Map columns using regex patterns.

    Args:
        sheet (Worksheet): The worksheet.
        header_row (int): 1-based header row number.
        registry (PatternRegistry): Pattern registry.
        sheet_type (Literal["invoice", "packing"]): Which column set to use.
        validation (ValidationResult): To record errors.

    Returns:
        ColumnMap | None: Column mapping, or None if required fields missing.
    """
    # Get column patterns based on sheet type
    if sheet_type == "invoice":
        columns = registry.invoice_columns
        required_fields = registry.get_required_invoice_fields()
    else:
        columns = registry.packing_columns
        required_fields = registry.get_required_packing_fields()

    # Read header row values
    headers: list[str] = []
    max_col = min(sheet.max_column or 20, 30)  # Limit scan range

    for col in range(1, max_col + 1):
        cell = sheet.cell(row=header_row, column=col)
        headers.append(normalize_header(cell.value))

    # Map columns using patterns
    column_map = ColumnMap(header_row=header_row)

    for field_name, field_pattern in columns.items():
        matched_col = _find_column(headers, field_pattern.patterns, field_name)
        if matched_col is not None:
            # Store as 0-based index
            column_map.columns[field_name] = matched_col

    # For invoice sheet: Try currency detection from data row if not found
    if sheet_type == "invoice" and "currency" not in column_map.columns:
        currency_col = _detect_currency_from_data(
            sheet, header_row, column_map.columns, headers
        )
        if currency_col is not None:
            column_map.columns["currency"] = currency_col

    # For invoice sheet: Adjust price/amount columns if they contain currency codes
    if sheet_type == "invoice":
        _adjust_numeric_columns_for_currency(sheet, header_row, column_map.columns, headers)

    # Check for missing required fields
    missing = [f for f in required_fields if f not in column_map.columns]
    if missing:
        for field in missing:
            validation.add_error(
                ErrorCode.REQUIRED_COLUMN_MISSING,
                f"Required column(s) missing: {field}",
            )
        return None

    return column_map


def map_columns_with_subheader(
    sheet: Worksheet,
    header_row: int,
    registry: PatternRegistry,
    sheet_type: Literal["invoice", "packing"],
    validation: ValidationResult,
) -> ColumnMap | None:
    """Map columns with support for multi-row headers.

    First tries header_row, then checks header_row + 1 for sub-headers
    if any required fields are missing.

    Args:
        sheet (Worksheet): The worksheet.
        header_row (int): 1-based header row number.
        registry (PatternRegistry): Pattern registry.
        sheet_type (Literal["invoice", "packing"]): Which column set to use.
        validation (ValidationResult): To record errors.

    Returns:
        ColumnMap | None: Column mapping, or None if required fields missing.
    """
    # Get column patterns based on sheet type
    if sheet_type == "invoice":
        columns = registry.invoice_columns
        required_fields = registry.get_required_invoice_fields()
    else:
        columns = registry.packing_columns
        required_fields = registry.get_required_packing_fields()

    # Read header row values
    headers: list[str] = []
    max_col = min(sheet.max_column or 20, 30)

    for col in range(1, max_col + 1):
        cell = sheet.cell(row=header_row, column=col)
        headers.append(normalize_header(cell.value))

    # First pass: map columns from header row
    column_map = ColumnMap(header_row=header_row)

    for field_name, field_pattern in columns.items():
        matched_col = _find_column(headers, field_pattern.patterns, field_name)
        if matched_col is not None:
            column_map.columns[field_name] = matched_col

    # Check which required fields are missing
    missing = [f for f in required_fields if f not in column_map.columns]

    # If any required fields missing, try sub-header row
    if missing:
        sub_row = header_row + 1
        sub_headers: list[str] = []

        for col in range(1, max_col + 1):
            cell = sheet.cell(row=sub_row, column=col)
            sub_headers.append(normalize_header(cell.value))

        # Check if this is actually a sub-header row (not data)
        # A row is a sub-header if it has descriptive text, not data values
        found_sub = False
        for field_name in missing:
            field_pattern = columns.get(field_name)
            if field_pattern is None:
                continue

            for col_idx, sub_value in enumerate(sub_headers):
                if not sub_value:
                    continue

                # Skip if this looks like a currency code (data value)
                if sub_value.upper() in CURRENCY_CODES:
                    continue

                # Try to match pattern
                for pattern in field_pattern.patterns:
                    if pattern.search(sub_value):
                        column_map.columns[field_name] = col_idx
                        found_sub = True
                        break
                if field_name in column_map.columns:
                    break

        if found_sub:
            column_map.has_subheader = True

    # For invoice sheet: Try currency detection from data row if not found
    if sheet_type == "invoice" and "currency" not in column_map.columns:
        data_row = header_row + (2 if column_map.has_subheader else 1)
        currency_col = _detect_currency_from_data(
            sheet, data_row - 1, column_map.columns, headers
        )
        if currency_col is not None:
            column_map.columns["currency"] = currency_col

    # For invoice sheet: Adjust price/amount columns if they contain currency codes
    if sheet_type == "invoice":
        data_row = header_row + (2 if column_map.has_subheader else 1)
        _adjust_numeric_columns_for_currency(
            sheet, data_row - 1, column_map.columns, headers
        )

    # Final check for missing required fields
    missing = [f for f in required_fields if f not in column_map.columns]
    if missing:
        for field in missing:
            validation.add_error(
                ErrorCode.REQUIRED_COLUMN_MISSING,
                f"Required column(s) missing: {field}",
            )
        return None

    return column_map


def _find_column(
    headers: list[str],
    patterns: list[re.Pattern[str]],
    field_name: str,
) -> int | None:
    """Find column index matching patterns.

    Multiple patterns matching same column: first pattern wins.
    Same pattern matching multiple columns: leftmost column wins.

    Args:
        headers (list[str]): Normalized header values.
        patterns (list[re.Pattern[str]]): Patterns to match.
        field_name (str): Field name for debugging.

    Returns:
        int | None: 0-based column index, or None if not found.
    """
    for pattern in patterns:
        for col_idx, header in enumerate(headers):
            if not header:
                continue
            if pattern.search(header):
                log_debug(
                    f"Column '{field_name}' matched at col {col_idx + 1} "
                    f"(header: '{header}', pattern: '{pattern.pattern}')"
                )
                return col_idx
    return None


def _detect_currency_from_data(
    sheet: Worksheet,
    header_row: int,
    mapped_columns: dict[str, int],
    headers: list[str],
) -> int | None:
    """Detect currency column from first data row values.

    Used when currency column header is not found via pattern matching.
    Scans first effective data row for currency codes.

    Per FR8: Columns matched to PRICE/AMOUNT headers are NOT skipped,
    as merged headers may have currency code in one sub-column.

    Args:
        sheet (Worksheet): The worksheet.
        header_row (int): 1-based header row number.
        mapped_columns (dict[str, int]): Already mapped columns (0-based).
        headers (list[str]): Header values.

    Returns:
        int | None: 0-based column index of currency column, or None.
    """
    # Find first non-blank data row
    data_row = _find_first_data_row(sheet, header_row + 1)
    if data_row is None:
        return None

    # Build set of columns to skip (already mapped, except price/amount)
    skip_cols = set()
    for field_name, col_idx in mapped_columns.items():
        # Don't skip if header contains price/amount keywords
        header_text = headers[col_idx].upper() if col_idx < len(headers) else ""
        if any(kw in header_text for kw in PRICE_AMOUNT_KEYWORDS):
            continue
        skip_cols.add(col_idx)

    # Scan data row for currency codes
    max_col = min(sheet.max_column or 20, 20)

    for col in range(1, max_col + 1):
        col_idx = col - 1

        if col_idx in skip_cols:
            continue

        cell = sheet.cell(row=data_row, column=col)
        value = str(cell.value).strip().upper() if cell.value else ""

        if value in CURRENCY_CODES:
            log_debug(f"Currency detected from data row: '{value}' at col {col}")
            return col_idx

    return None


def _adjust_numeric_columns_for_currency(
    sheet: Worksheet,
    header_row: int,
    mapped_columns: dict[str, int],
    headers: list[str],
) -> None:
    """Adjust price/amount column mapping when they contain currency codes.

    When merged headers are used, the header pattern may match a column
    containing currency code, while the actual numeric value is in col+1.

    Args:
        sheet (Worksheet): The worksheet.
        header_row (int): 1-based header row number.
        mapped_columns (dict[str, int]): Column mapping to adjust (0-based).
        headers (list[str]): Header values.
    """
    data_row = _find_first_data_row(sheet, header_row + 1)
    if data_row is None:
        return

    for field_name in ["price", "amount"]:
        if field_name not in mapped_columns:
            continue

        col_idx = mapped_columns[field_name]
        cell = sheet.cell(row=data_row, column=col_idx + 1)
        value = str(cell.value).strip().upper() if cell.value else ""

        # Check if current column has currency code
        if value in CURRENCY_CODES:
            # Check if next column has numeric value
            next_col = col_idx + 1
            if next_col < (sheet.max_column or 20):
                next_cell = sheet.cell(row=data_row, column=next_col + 1)
                next_value = next_cell.value

                if next_value is not None and _is_numeric(next_value):
                    log_debug(
                        f"Adjusting {field_name} column from {col_idx + 1} to {next_col + 1} "
                        f"(currency in original, numeric in adjacent)"
                    )
                    mapped_columns[field_name] = next_col


def _find_first_data_row(sheet: Worksheet, start_row: int) -> int | None:
    """Find first non-blank data row.

    Args:
        sheet (Worksheet): The worksheet.
        start_row (int): 1-based row to start searching.

    Returns:
        int | None: 1-based row number, or None if not found.
    """
    max_row = min(sheet.max_row or 100, start_row + 10)

    for row in range(start_row, max_row + 1):
        non_empty = 0
        for col in range(1, 11):
            cell = sheet.cell(row=row, column=col)
            if cell.value is not None:
                non_empty += 1

        if non_empty >= 3:
            return row

    return None


def _is_numeric(value: object) -> bool:
    """Check if a value is numeric.

    Args:
        value: Value to check.

    Returns:
        bool: True if numeric.
    """
    if isinstance(value, (int, float)):
        return True

    try:
        float(str(value).replace(",", ""))
        return True
    except (ValueError, TypeError):
        return False
