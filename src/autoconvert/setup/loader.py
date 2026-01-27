"""Configuration file loaders for YAML and Excel files."""

import re
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

from autoconvert.core.errors import ConfigError, ErrorCode


def load_yaml_config(config_path: Path) -> dict[str, Any]:
    """Load and parse YAML configuration file.

    Args:
        config_path (Path): Path to field_patterns.yaml.

    Returns:
        dict[str, Any]: Parsed YAML content.

    Raises:
        ConfigError: If file not found or malformed.
    """
    if not config_path.exists():
        raise ConfigError(
            ErrorCode.CONFIG_NOT_FOUND,
            f"Configuration file not found: {config_path}",
        )

    try:
        with open(config_path, "r", encoding="utf-8") as f:
            config = yaml.safe_load(f)
    except yaml.YAMLError as e:
        raise ConfigError(
            ErrorCode.MALFORMED_CONFIG,
            f"Invalid YAML syntax in {config_path}: {e}",
        )
    except Exception as e:
        raise ConfigError(
            ErrorCode.MALFORMED_CONFIG,
            f"Failed to read {config_path}: {e}",
        )

    if not isinstance(config, dict):
        raise ConfigError(
            ErrorCode.MALFORMED_CONFIG,
            f"Configuration must be a YAML dictionary, got {type(config).__name__}",
        )

    return config


def validate_yaml_structure(config: dict[str, Any]) -> None:
    """Validate required root keys in YAML configuration.

    Args:
        config (dict[str, Any]): Parsed YAML configuration.

    Raises:
        ConfigError: If required keys are missing.
    """
    required_keys = [
        "invoice_sheet",
        "packing_sheet",
        "invoice_columns",
        "packing_columns",
        "inv_no_cell",
    ]

    missing = [key for key in required_keys if key not in config]
    if missing:
        raise ConfigError(
            ErrorCode.MALFORMED_CONFIG,
            f"Missing required config keys: {', '.join(missing)}",
        )


def compile_patterns(patterns: list[str], context: str) -> list[re.Pattern[str]]:
    """Compile regex patterns with validation.

    Args:
        patterns (list[str]): List of regex pattern strings.
        context (str): Description for error messages.

    Returns:
        list[re.Pattern[str]]: Compiled regex patterns.

    Raises:
        ConfigError: If any pattern has invalid syntax.
    """
    compiled = []
    for pattern in patterns:
        try:
            compiled.append(re.compile(pattern))
        except re.error as e:
            raise ConfigError(
                ErrorCode.INVALID_REGEX,
                f"Invalid regex in {context}: '{pattern}' - {e}",
            )
    return compiled


def load_lookup_table(excel_path: Path) -> dict[str, str]:
    """Load lookup table from Excel file.

    Expects two columns: A = Source_Value, B = Target_Code
    Header in row 1, data starts row 2.

    Args:
        excel_path (Path): Path to lookup Excel file.

    Returns:
        dict[str, str]: Mapping of source values to target codes.

    Raises:
        ConfigError: If file not found, corrupted, or has duplicates.
    """
    if not excel_path.exists():
        raise ConfigError(
            ErrorCode.CONFIG_NOT_FOUND,
            f"Lookup file not found: {excel_path}",
        )

    try:
        wb = load_workbook(excel_path, read_only=True, data_only=True)
        sheet = wb.active
    except Exception as e:
        raise ConfigError(
            ErrorCode.MALFORMED_CONFIG,
            f"Failed to open lookup file {excel_path}: {e}",
        )

    if sheet is None:
        raise ConfigError(
            ErrorCode.MALFORMED_CONFIG,
            f"Lookup file has no active sheet: {excel_path}",
        )

    lookup = {}
    seen_sources: dict[str, int] = {}  # Track source values and their row numbers

    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if len(row) < 2:
            continue

        source_val = row[0]
        target_code = row[1]

        if source_val is None or target_code is None:
            continue

        # Normalize source value: uppercase, trimmed
        source_key = str(source_val).strip().upper()
        target_str = str(target_code).strip()

        if not source_key:
            continue

        # Check for duplicates
        if source_key in seen_sources:
            raise ConfigError(
                ErrorCode.DUPLICATE_LOOKUP,
                f"Duplicate source value '{source_key}' in {excel_path.name} "
                f"(rows {seen_sources[source_key]} and {row_num})",
            )

        seen_sources[source_key] = row_num
        lookup[source_key] = target_str

    wb.close()
    return lookup


def validate_template(template_path: Path) -> None:
    """Validate output template structure.

    Args:
        template_path (Path): Path to output_template.xlsx.

    Raises:
        ConfigError: If template not found or has wrong structure.
    """
    if not template_path.exists():
        raise ConfigError(
            ErrorCode.CONFIG_NOT_FOUND,
            f"Output template not found: {template_path}",
        )

    try:
        wb = load_workbook(template_path, read_only=True)
        sheet = wb.active
    except Exception as e:
        raise ConfigError(
            ErrorCode.MALFORMED_CONFIG,
            f"Failed to open template file {template_path}: {e}",
        )

    if sheet is None:
        raise ConfigError(
            ErrorCode.TEMPLATE_INVALID,
            "Template file has no active sheet",
        )

    # Check for 40 columns (A-AN)
    max_col = sheet.max_column
    if max_col is None or max_col < 40:
        raise ConfigError(
            ErrorCode.TEMPLATE_INVALID,
            f"Template must have at least 40 columns (A-AN), found {max_col}",
        )

    wb.close()
