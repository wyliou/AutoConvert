"""Invoice sheet data extraction."""

import re

from openpyxl.worksheet.worksheet import Worksheet

from autoconvert.core.errors import ErrorCode
from autoconvert.core.models import ColumnMap, InvoiceItem, ValidationResult
from autoconvert.logging.messages import log_debug, log_extraction, log_inv_no_extracted
from autoconvert.parsing.merged_cells import MergeTracker
from autoconvert.setup.registry import PatternRegistry
from autoconvert.transformation.cleaners import clean_invoice_number

# Stop keywords for data extraction
STOP_KEYWORDS = {"total", "合计", "总计", "小计"}
FOOTER_KEYWORDS = {"报关行", "有限公司"}

# Unit suffixes to strip from numeric values
UNIT_SUFFIXES = re.compile(r"(KGS?|PCS|EA|件|个)$", re.IGNORECASE)

# Labels to remove when checking for pure label
LABEL_KEYWORDS = {
    "invoice",
    "inv",
    "no",
    "number",
    "date",
    "&",
    "of",
    ".",
    ":",
    "：",
}


def extract_invoice_items(
    sheet: Worksheet,
    column_map: ColumnMap,
    tracker: MergeTracker,
    registry: PatternRegistry,
    validation: ValidationResult,
) -> list[InvoiceItem]:
    """Extract invoice line items from sheet.

    Args:
        sheet (Worksheet): The invoice sheet.
        column_map (ColumnMap): Column mapping (0-based indices).
        tracker (MergeTracker): Merge tracker for handling merged cells.
        registry (PatternRegistry): Pattern registry.
        validation (ValidationResult): For recording errors.

    Returns:
        list[InvoiceItem]: Extracted invoice items.
    """
    items: list[InvoiceItem] = []
    header_row = column_map.header_row
    start_row = header_row + (2 if column_map.has_subheader else 1)

    # Get column indices (convert from 0-based to field access)
    cols = column_map.columns

    # Find string columns for merge propagation
    string_fields = [
        "part_no",
        "po_no",
        "currency",
        "coo",
        "cod",
        "brand",
        "brand_type",
        "model",
        "serial",
        "inv_no",
    ]
    string_col_indices = [cols[f] + 1 for f in string_fields if f in cols]

    max_row = sheet.max_row or 1000
    first_data_found = False

    for row_num in range(start_row, max_row + 1):
        # Check stop conditions first (check first 10 columns)
        if _is_stop_row(sheet, row_num, cols):
            log_debug(f"Stop condition at row {row_num}")
            break

        # Check if row is blank
        if _is_blank_row(sheet, row_num, cols):
            if not first_data_found:
                continue  # Skip leading blank rows
            else:
                break  # Stop at trailing blank rows

        # Extract row data
        item = _extract_row(sheet, row_num, cols, tracker)

        # Skip header continuation rows
        if item.part_no and "part no" in item.part_no.lower():
            continue

        # Check for stop condition based on data
        if first_data_found:
            if not item.part_no and item.qty == 0:
                break

        if item.part_no or item.qty > 0:
            items.append(item)
            first_data_found = True

    # Log extraction results
    if items:
        log_extraction(
            "Invoice sheet",
            len(items),
            items[0].row,
            items[-1].row,
        )

    return items


def _extract_row(
    sheet: Worksheet,
    row_num: int,
    cols: dict[str, int],
    tracker: MergeTracker,
) -> InvoiceItem:
    """Extract a single invoice item from a row.

    Args:
        sheet (Worksheet): The worksheet.
        row_num (int): 1-based row number.
        cols (dict[str, int]): Column mapping (0-based indices).
        tracker (MergeTracker): Merge tracker.

    Returns:
        InvoiceItem: Extracted item.
    """
    item = InvoiceItem(row=row_num)

    # Extract string fields
    for field in ["part_no", "po_no", "currency", "coo", "cod", "brand", "brand_type", "model", "serial", "inv_no"]:
        if field in cols:
            value = _get_string_value(sheet, row_num, cols[field] + 1, tracker)
            setattr(item, field, value)

    # Extract numeric fields with precision handling
    if "qty" in cols:
        item.qty = _get_numeric_value(sheet, row_num, cols["qty"] + 1)

    if "price" in cols:
        item.price = _get_numeric_value(sheet, row_num, cols["price"] + 1, precision=5)

    if "amount" in cols:
        item.amount = _get_numeric_value(sheet, row_num, cols["amount"] + 1, precision=2)

    return item


def _get_string_value(
    sheet: Worksheet,
    row: int,
    col: int,
    tracker: MergeTracker,
) -> str:
    """Get string value from cell, handling merged cells.

    Args:
        sheet (Worksheet): The worksheet.
        row (int): 1-based row number.
        col (int): 1-based column number.
        tracker (MergeTracker): Merge tracker.

    Returns:
        str: String value, empty string if None.
    """
    # Check if cell was part of a merge and get origin value
    origin = tracker.get_origin(row, col)
    if origin is not None:
        cell = sheet.cell(row=origin[0], column=origin[1])
    else:
        cell = sheet.cell(row=row, column=col)

    value = cell.value
    if value is None:
        return ""

    return str(value).strip()


def _get_numeric_value(
    sheet: Worksheet,
    row: int,
    col: int,
    precision: int | None = None,
) -> float:
    """Get numeric value from cell with optional precision.

    Args:
        sheet (Worksheet): The worksheet.
        row (int): 1-based row number.
        col (int): 1-based column number.
        precision (int | None): Decimal precision to round to.

    Returns:
        float: Numeric value, 0.0 if not numeric.
    """
    cell = sheet.cell(row=row, column=col)
    value = cell.value

    if value is None:
        return 0.0

    if isinstance(value, (int, float)):
        num = float(value)
    else:
        # Try to parse string
        str_val = str(value).strip()
        # Strip unit suffixes
        str_val = UNIT_SUFFIXES.sub("", str_val).strip()
        # Remove commas
        str_val = str_val.replace(",", "")

        try:
            num = float(str_val)
        except (ValueError, TypeError):
            return 0.0

    # Apply precision if specified
    if precision is not None:
        # Use ROUND_HALF_UP via epsilon trick
        factor = 10**precision
        num = round(num * factor + 1e-9) / factor

    return num


def _is_stop_row(sheet: Worksheet, row_num: int, cols: dict[str, int]) -> bool:
    """Check if row contains stop keywords.

    Args:
        sheet (Worksheet): The worksheet.
        row_num (int): 1-based row number.
        cols (dict[str, int]): Column mapping.

    Returns:
        bool: True if stop condition detected.
    """
    # Check first 10 columns for stop keywords
    for col in range(1, 11):
        cell = sheet.cell(row=row_num, column=col)
        value = cell.value

        if value is None:
            continue

        str_val = str(value).strip().lower()

        # Check for stop keywords
        for kw in STOP_KEYWORDS:
            if kw in str_val:
                return True

        # Check for footer keywords
        for kw in FOOTER_KEYWORDS:
            if kw in str_val:
                return True

    # Check part_no for "total"
    if "part_no" in cols:
        part_cell = sheet.cell(row=row_num, column=cols["part_no"] + 1)
        if part_cell.value:
            part_val = str(part_cell.value).strip().lower()
            if "total" in part_val:
                return True

    return False


def _is_blank_row(sheet: Worksheet, row_num: int, cols: dict[str, int]) -> bool:
    """Check if row is blank (all key columns empty).

    Args:
        sheet (Worksheet): The worksheet.
        row_num (int): 1-based row number.
        cols (dict[str, int]): Column mapping.

    Returns:
        bool: True if row is blank.
    """
    # Check key columns
    key_fields = ["part_no", "qty", "price", "amount"]

    for field in key_fields:
        if field in cols:
            cell = sheet.cell(row=row_num, column=cols[field] + 1)
            if cell.value is not None and str(cell.value).strip():
                return False

    return True


def extract_inv_no_from_header(
    sheet: Worksheet,
    registry: PatternRegistry,
) -> tuple[str | None, str | None, str | None]:
    """Extract invoice number from header area (rows 2-15).

    Uses two methods:
    1. Label match - find label pattern, then check adjacent cells
    2. Embedded value - extract from cell containing both label and value

    Args:
        sheet (Worksheet): The invoice sheet.
        registry (PatternRegistry): Pattern registry with inv_no patterns.

    Returns:
        tuple[str | None, str | None, str | None]:
            (inv_no, method_description, label_found)
    """
    # Method 1: Label matching
    for row in range(2, 16):
        for col in range(1, 12):
            cell = sheet.cell(row=row, column=col)
            value = cell.value

            if value is None:
                continue

            str_val = str(value).strip()

            # Check label patterns
            for pattern in registry.inv_no_label_patterns:
                if pattern.search(str_val):
                    # Found a label, check adjacent cells
                    result = _search_adjacent_for_inv_no(
                        sheet, row, col, str_val, registry, depth=0
                    )
                    if result:
                        inv_no, method, cell_ref = result
                        # Clean the invoice number
                        inv_no = clean_invoice_number(inv_no)
                        log_inv_no_extracted(method, str_val, inv_no, cell_ref)
                        return inv_no, method, str_val

    # Method 2: Embedded value
    for row in range(2, 16):
        for col in range(1, 12):
            cell = sheet.cell(row=row, column=col)
            value = cell.value

            if value is None:
                continue

            str_val = str(value).strip()

            # Check embedded patterns
            for pattern in registry.inv_no_patterns:
                match = pattern.search(str_val)
                if match:
                    inv_no = match.group(1) if match.groups() else match.group(0)

                    # Check exclude patterns
                    if not _is_excluded(inv_no, registry.inv_no_exclude_patterns):
                        inv_no = clean_invoice_number(inv_no)
                        cell_ref = f"{_col_letter(col)}{row}"
                        log_inv_no_extracted("embedded", str_val, inv_no, cell_ref)
                        return inv_no, "embedded", str_val

    return None, None, None


def _search_adjacent_for_inv_no(
    sheet: Worksheet,
    row: int,
    col: int,
    label: str,
    registry: PatternRegistry,
    depth: int,
) -> tuple[str, str, str] | None:
    """Search adjacent cells for invoice number.

    Args:
        sheet (Worksheet): The worksheet.
        row (int): Label row (1-based).
        col (int): Label column (1-based).
        label (str): The label text found.
        registry (PatternRegistry): Pattern registry.
        depth (int): Current recursion depth.

    Returns:
        tuple[str, str, str] | None: (inv_no, method, cell_ref) or None.
    """
    if depth > 3:
        return None

    # Adjacent positions to check: right, right+1, below, below+1
    positions = [
        (row, col + 1, "1 cell right"),
        (row, col + 2, "2 cells right"),
        (row + 1, col, "1 row below"),
        (row + 2, col, "2 rows below"),
    ]

    for adj_row, adj_col, method in positions:
        if adj_row > 15 or adj_col > 15:
            continue

        cell = sheet.cell(row=adj_row, column=adj_col)
        value = cell.value

        if value is None:
            continue

        str_val = str(value).strip()

        if not str_val:
            continue

        # Check if this is a pure label (nested label)
        if _is_pure_label(str_val, registry.inv_no_label_patterns):
            # Recursive search from this label
            result = _search_adjacent_for_inv_no(
                sheet, adj_row, adj_col, str_val, registry, depth + 1
            )
            if result:
                inv_no, _, cell_ref = result
                return inv_no, f"nested label of '{str_val}'", cell_ref
            continue

        # Check exclude patterns
        if _is_excluded(str_val, registry.inv_no_exclude_patterns):
            # Try extended search
            extended_positions = [
                (adj_row, adj_col + 1, "1 cell right"),
                (adj_row + 1, adj_col, "1 row below"),
            ]
            for ext_row, ext_col, ext_method in extended_positions:
                ext_cell = sheet.cell(row=ext_row, column=ext_col)
                ext_value = ext_cell.value
                if ext_value:
                    ext_str = str(ext_value).strip()
                    if ext_str and not _is_excluded(ext_str, registry.inv_no_exclude_patterns):
                        if not _is_pure_label(ext_str, registry.inv_no_label_patterns):
                            cell_ref = f"{_col_letter(ext_col)}{ext_row}"
                            return ext_str, method, cell_ref
            continue

        # Valid value found
        cell_ref = f"{_col_letter(adj_col)}{adj_row}"
        return str_val, method, cell_ref

    return None


def _is_pure_label(value: str, label_patterns: list[re.Pattern[str]]) -> bool:
    """Check if value is a pure label (not containing an invoice number).

    A pure label matches a label pattern AND has <3 alphanumeric chars
    after removing label keywords.

    Args:
        value (str): Value to check.
        label_patterns (list[re.Pattern[str]]): Label patterns.

    Returns:
        bool: True if this is a pure label.
    """
    if not any(pattern.search(value) for pattern in label_patterns):
        return False

    # Remove keywords in length-descending order (CRITICAL per FR15)
    remaining = value.lower()
    sorted_keywords = sorted(LABEL_KEYWORDS, key=len, reverse=True)

    for kw in sorted_keywords:
        remaining = remaining.replace(kw.lower(), "")

    # Remove punctuation and whitespace
    remaining = re.sub(r"[^\w]", "", remaining)

    # Pure label if <3 alphanumeric chars remain
    return len(remaining) < 3


def _is_excluded(value: str, exclude_patterns: list[re.Pattern[str]]) -> bool:
    """Check if value matches any exclude pattern.

    Args:
        value (str): Value to check.
        exclude_patterns (list[re.Pattern[str]]): Patterns to exclude.

    Returns:
        bool: True if value should be excluded.
    """
    return any(pattern.search(value) for pattern in exclude_patterns)


def _col_letter(col: int) -> str:
    """Convert 1-based column number to letter.

    Args:
        col (int): 1-based column number.

    Returns:
        str: Column letter (A, B, ..., Z, AA, AB, ...).
    """
    result = ""
    while col > 0:
        col, remainder = divmod(col - 1, 26)
        result = chr(65 + remainder) + result
    return result
