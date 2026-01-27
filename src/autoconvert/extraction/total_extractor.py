"""Total value extraction from packing sheet."""

import re
from decimal import ROUND_HALF_UP, Decimal

from openpyxl.worksheet.worksheet import Worksheet

from autoconvert.core.errors import ErrorCode, WarningCode
from autoconvert.core.models import ColumnMap, PackingTotals, ValidationResult
from autoconvert.logging.messages import log_debug, log_packing_totals

# Unit suffixes to strip from weight values
WEIGHT_UNITS = re.compile(r"(KGS?|G|LBS?)$", re.IGNORECASE)

# Patterns for packet extraction
PACKET_LABELS = {"件数", "件數"}
PLT_PATTERN = re.compile(r"^(\d+)\s*PLT\.?\s*G?$", re.IGNORECASE)
PLT_REVERSE_PATTERN = re.compile(r"^PLT\.?\s*G", re.IGNORECASE)
UNIT_SUFFIX_PATTERN = re.compile(r"^(\d+)\s*(CTNS?|箱|托|件|PCS)", re.IGNORECASE)
TOTAL_BREAKDOWN_PATTERN = re.compile(r"^(\d+)\s*[（(]")  # 348（256胶框+92纸箱）
EMBEDDED_CHINESE_PATTERN = re.compile(r"共?(\d+)\s*托")
PALLET_RANGE_PATTERN = re.compile(r"PLT#(\d+)\s*\(", re.IGNORECASE)  # PLT#1(1~34) → 1


def extract_packing_totals(
    sheet: Worksheet,
    column_map: ColumnMap,
    total_row: int,
    validation: ValidationResult,
) -> PackingTotals | None:
    """Extract total values from packing sheet total row.

    Args:
        sheet (Worksheet): The packing sheet.
        column_map (ColumnMap): Column mapping.
        total_row (int): 1-based total row number.
        validation (ValidationResult): For recording errors/warnings.

    Returns:
        PackingTotals | None: Extracted totals, or None if critical values missing.
    """
    cols = column_map.columns
    totals = PackingTotals(total_row=total_row)

    # Extract total_nw
    if "nw" in cols:
        nw_col = cols["nw"] + 1
        nw_value, nw_precision = _extract_weight(sheet, total_row, nw_col)

        if nw_value is None or nw_value <= 0:
            validation.add_error(
                ErrorCode.INVALID_TOTAL_NW,
                "total_nw is 0, negative, or non-numeric",
            )
            return None

        totals.total_nw = nw_value
        totals.nw_precision = nw_precision
    else:
        validation.add_error(
            ErrorCode.INVALID_TOTAL_NW,
            "NW column not found",
        )
        return None

    # Extract total_gw (with packaging weight check)
    if "gw" in cols:
        gw_col = cols["gw"] + 1
        gw_value, gw_precision = _extract_gw_with_packaging(sheet, total_row, gw_col)

        if gw_value is None or gw_value <= 0:
            validation.add_error(
                ErrorCode.INVALID_TOTAL_GW,
                "total_gw is 0, negative, or non-numeric",
            )
            return None

        totals.total_gw = gw_value
        totals.gw_precision = gw_precision
    else:
        validation.add_error(
            ErrorCode.INVALID_TOTAL_GW,
            "GW column not found",
        )
        return None

    # Extract total_packets (optional)
    nw_col = cols.get("nw", 10)  # Default to column 10 if not found
    packets = _extract_packets(sheet, total_row, nw_col + 1)

    if packets is not None:
        totals.total_packets = packets
    else:
        validation.add_warning(
            WarningCode.MISSING_TOTAL_PACKETS,
            "total_packets not found or invalid, please verify manually in output",
        )

    # Log totals
    nw_str = _format_decimal(totals.total_nw)
    gw_str = _format_decimal(totals.total_gw)
    log_packing_totals(total_row, nw_str, gw_str, totals.total_packets)

    return totals


def _extract_weight(
    sheet: Worksheet,
    row: int,
    col: int,
) -> tuple[Decimal | None, int]:
    """Extract weight value with precision detection.

    Args:
        sheet (Worksheet): The worksheet.
        row (int): 1-based row number.
        col (int): 1-based column number.

    Returns:
        tuple[Decimal | None, int]: (weight value, decimal precision).
    """
    cell = sheet.cell(row=row, column=col)
    value = cell.value
    number_format = cell.number_format or "General"

    if value is None:
        return None, 2

    # Handle string values (may have unit suffix)
    if isinstance(value, str):
        str_val = value.strip()
        # Strip unit suffix BEFORE precision detection
        str_val_clean = WEIGHT_UNITS.sub("", str_val).strip()
        str_val_clean = str_val_clean.replace(",", "")

        try:
            num = float(str_val_clean)
        except ValueError:
            return None, 2

        # Detect precision from cleaned string value
        precision = _detect_precision_from_string(str_val_clean)

        # Clean floating-point artifacts
        num = round(num, 5)
        dec_val = Decimal(str(num)).quantize(
            Decimal(10) ** -precision,
            rounding=ROUND_HALF_UP,
        )
        return dec_val, precision

    # Handle numeric values
    if isinstance(value, (int, float)):
        num = float(value)

        # Detect precision from number format
        precision = _detect_precision_from_format(number_format)

        if precision is None:
            # No format, clean artifacts and detect from value
            num = round(num, 5)
            precision = _detect_precision_from_value(num)

        dec_val = Decimal(str(num)).quantize(
            Decimal(10) ** -precision,
            rounding=ROUND_HALF_UP,
        )
        return dec_val, precision

    return None, 2


def _extract_gw_with_packaging(
    sheet: Worksheet,
    total_row: int,
    gw_col: int,
) -> tuple[Decimal | None, int]:
    """Extract GW with packaging weight check.

    Some vendors add packaging weight below total row. If rows +1 and +2
    both have numeric values, use +2 as final total_gw.

    Args:
        sheet (Worksheet): The worksheet.
        total_row (int): 1-based total row number.
        gw_col (int): 1-based GW column number.

    Returns:
        tuple[Decimal | None, int]: (GW value, decimal precision).
    """
    # First get value from total row
    gw_value, gw_precision = _extract_weight(sheet, total_row, gw_col)

    # Check +1 and +2 rows for packaging weight pattern
    row_plus1 = sheet.cell(row=total_row + 1, column=gw_col).value
    row_plus2 = sheet.cell(row=total_row + 2, column=gw_col).value

    if row_plus1 is not None and row_plus2 is not None:
        try:
            val1 = float(str(row_plus1).replace(",", "").strip())
            val2 = float(str(row_plus2).replace(",", "").strip())

            # Both have numeric values - use +2 as final GW
            if val1 > 0 and val2 > 0:
                log_debug(
                    f"Packaging weight detected: total row GW={gw_value}, "
                    f"row+1={val1}, row+2={val2}, using {val2}"
                )
                return _extract_weight(sheet, total_row + 2, gw_col)

        except (ValueError, TypeError):
            pass

    return gw_value, gw_precision


def _extract_packets(
    sheet: Worksheet,
    total_row: int,
    nw_col: int,
) -> int | None:
    """Extract total packets using multi-priority search.

    Priority 1: Packet/Carton labels (件数, 件數)
    Priority 2: PLT.G indicator
    Priority 3: Below total row patterns

    Args:
        sheet (Worksheet): The worksheet.
        total_row (int): 1-based total row number.
        nw_col (int): 1-based NW column number for search range.

    Returns:
        int | None: Packet count, or None if not found.
    """
    max_col = max(nw_col + 2, 11)

    # Priority 1: Packet labels in rows below total
    for row_offset in range(1, 4):
        row = total_row + row_offset

        for col in range(1, max_col + 1):
            cell = sheet.cell(row=row, column=col)
            value = cell.value

            if value is None:
                continue

            str_val = str(value).strip()

            # Check for packet labels
            for label in PACKET_LABELS:
                if label in str_val:
                    # Look for value adjacent to label
                    packets = _search_adjacent_for_packets(sheet, row, col, max_col)
                    if packets is not None:
                        return packets

                    # Try embedded value
                    match = re.search(r"[:：]\s*(\d+)", str_val)
                    if match:
                        return int(match.group(1))

    # Priority 2: PLT.G indicator (above total row)
    for row_offset in range(1, 3):
        row = total_row - row_offset

        for col in range(1, max_col + 1):
            cell = sheet.cell(row=row, column=col)
            value = cell.value

            if value is None:
                continue

            str_val = str(value).strip()

            # Number-before-PLT format: "7 PLT.G"
            match = PLT_PATTERN.match(str_val)
            if match:
                return int(match.group(1))

            # PLT-before-number format: "PLT.G 5"
            if PLT_REVERSE_PATTERN.match(str_val):
                # Check adjacent cell
                adj_cell = sheet.cell(row=row, column=col + 1)
                if adj_cell.value:
                    try:
                        return int(str(adj_cell.value).strip())
                    except ValueError:
                        pass

    # Priority 3: Below total row patterns
    for row_offset in range(1, 4):
        row = total_row + row_offset

        for col in range(1, max_col + 1):
            cell = sheet.cell(row=row, column=col)
            value = cell.value

            if value is None:
                continue

            str_val = str(value).strip()

            # Total with breakdown pattern: "348（256胶框+92纸箱）"
            match = TOTAL_BREAKDOWN_PATTERN.match(str_val)
            if match:
                return int(match.group(1))

            # Unit suffix pattern: "7CTNS", "30箱", "50件"
            match = UNIT_SUFFIX_PATTERN.match(str_val)
            if match:
                return int(match.group(1))

            # Embedded Chinese pattern: "共7托"
            match = EMBEDDED_CHINESE_PATTERN.search(str_val)
            if match:
                return int(match.group(1))

            # Pallet range pattern: "PLT#1(1~34)" → extract pallet count (1)
            match = PALLET_RANGE_PATTERN.search(str_val)
            if match:
                return int(match.group(1))

    return None


def _search_adjacent_for_packets(
    sheet: Worksheet,
    row: int,
    col: int,
    max_col: int,
) -> int | None:
    """Search adjacent cells for packet value.

    Args:
        sheet (Worksheet): The worksheet.
        row (int): 1-based row of label.
        col (int): 1-based column of label.
        max_col (int): Maximum column to search.

    Returns:
        int | None: Packet count, or None.
    """
    # Check cells to the right
    for offset in range(1, 4):
        adj_col = col + offset
        if adj_col > max_col:
            break

        cell = sheet.cell(row=row, column=adj_col)
        value = cell.value

        if value is None:
            continue

        str_val = str(value).strip()

        # Try to extract number with optional unit
        match = re.match(r"^(\d+)", str_val)
        if match:
            packets = int(match.group(1))
            if 1 <= packets <= 1000:
                return packets

    return None


def _detect_precision_from_format(number_format: str) -> int | None:
    """Detect decimal precision from Excel number format.

    Args:
        number_format (str): Excel number format string.

    Returns:
        int | None: Decimal precision, or None if format is General.
    """
    if number_format == "General" or not number_format:
        return None

    # Find decimal point and count zeros after
    if "." not in number_format:
        return 0

    # Extract decimal part
    parts = number_format.split(".")
    if len(parts) < 2:
        return 0

    decimal_part = parts[1]

    # Count '0' placeholders
    precision = 0
    for char in decimal_part:
        if char == "0":
            precision += 1
        elif char in "#?,;_()[]/":
            continue
        else:
            break

    return precision


def _detect_precision_from_string(value: str) -> int:
    """Detect decimal precision from string value.

    Args:
        value (str): String representation of number.

    Returns:
        int: Decimal precision (minimum 2).
    """
    if "." not in value:
        return 2

    decimal_part = value.split(".")[1]
    return max(len(decimal_part.rstrip("0")), 2)


def _detect_precision_from_value(value: float) -> int:
    """Detect decimal precision from float value.

    Args:
        value (float): The numeric value.

    Returns:
        int: Decimal precision (minimum 2, maximum 5).
    """
    str_val = str(value)

    if "." not in str_val:
        return 2

    decimal_part = str_val.split(".")[1]

    # Remove trailing zeros for display precision
    precision = len(decimal_part.rstrip("0"))

    return max(min(precision, 5), 2)


def _format_decimal(value: Decimal) -> str:
    """Format Decimal for display, removing trailing zeros.

    Args:
        value (Decimal): The Decimal value.

    Returns:
        str: Formatted string.
    """
    # Normalize to remove trailing zeros
    normalized = value.normalize()

    # Ensure at least one decimal place for whole numbers
    str_val = str(normalized)
    if "." not in str_val:
        str_val += ".0"

    return str_val
