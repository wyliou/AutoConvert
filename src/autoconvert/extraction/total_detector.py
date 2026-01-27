"""Total row detection in packing sheet."""

from openpyxl.worksheet.worksheet import Worksheet

from autoconvert.core.models import ColumnMap
from autoconvert.logging.messages import log_debug
from autoconvert.parsing.merged_cells import MergeTracker

# Keywords for total row detection
TOTAL_KEYWORDS = {"total", "合计", "总计", "小计"}


def detect_total_row(
    sheet: Worksheet,
    column_map: ColumnMap,
    last_data_row: int,
    tracker: MergeTracker,
) -> int | None:
    """Detect total row using two-strategy approach.

    Strategy 1: Keyword matching - search for TOTAL keywords
    Strategy 2: Implicit total - empty part_no with numeric NW and GW

    Args:
        sheet (Worksheet): The packing sheet.
        column_map (ColumnMap): Column mapping.
        last_data_row (int): Last extracted packing data row (1-based).
        tracker (MergeTracker): Merge tracker.

    Returns:
        int | None: 1-based row number of total row, or None if not found.
    """
    cols = column_map.columns
    search_start = last_data_row + 1
    search_end = min(last_data_row + 15, sheet.max_row or last_data_row + 15)

    # Strategy 1: Keyword matching
    for row in range(search_start, search_end + 1):
        # Check first 10 columns for TOTAL keywords
        for col in range(1, 11):
            cell = sheet.cell(row=row, column=col)
            value = cell.value

            if value is None:
                continue

            str_val = str(value).strip().lower()

            for kw in TOTAL_KEYWORDS:
                if kw in str_val:
                    log_debug(f"Total row detected (keyword '{kw}') at row {row}")
                    return row

    # Strategy 2: Implicit total row (empty part_no, numeric NW and GW)
    if "part_no" in cols and "nw" in cols and "gw" in cols:
        part_col = cols["part_no"] + 1
        nw_col = cols["nw"] + 1
        gw_col = cols["gw"] + 1

        for row in range(search_start, search_end + 1):
            # Check if part_no is empty
            part_cell = sheet.cell(row=row, column=part_col)
            part_value = part_cell.value

            if part_value is not None and str(part_value).strip():
                continue  # Has part_no, not a total row

            # Check if this cell was part of a merge (exclude merged data rows)
            if tracker.is_merged(row, part_col):
                origin = tracker.get_origin(row, part_col)
                if origin and origin[0] < row:
                    # This is a continuation row from a merge, not total
                    continue

            # Check NW and GW for numeric values > 0
            nw_value = _get_numeric(sheet, row, nw_col)
            gw_value = _get_numeric(sheet, row, gw_col)

            if nw_value > 0 and gw_value > 0:
                log_debug(f"Total row detected (implicit) at row {row}")
                return row

    return None


def _get_numeric(sheet: Worksheet, row: int, col: int) -> float:
    """Get numeric value from cell.

    Args:
        sheet (Worksheet): The worksheet.
        row (int): 1-based row number.
        col (int): 1-based column number.

    Returns:
        float: Numeric value, 0.0 if not numeric.
    """
    cell = sheet.cell(row=row, column=col)
    value = cell.value

    if value is None:
        return 0.0

    if isinstance(value, (int, float)):
        return float(value)

    try:
        return float(str(value).replace(",", "").strip())
    except (ValueError, TypeError):
        return 0.0
